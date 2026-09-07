from __future__ import annotations

import io
import os
from datetime import UTC, datetime
from decimal import Decimal

import boto3
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill


def env(name: str) -> str:
    value = os.getenv(name)
    if not value:
        raise RuntimeError(f"{name} must be configured")
    return value


def excel_value(value):
    return float(value) if isinstance(value, Decimal) else value


def master_workbook(calls: list[dict]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Call Intelligence"
    headers = ["Call SID", "Status", "Processed", "Sentiment", "Performance", "Resolved", "Revenue Opportunity", "Keywords", "Objections", "Missed Opportunities", "Training Recommendations"]
    sheet.append(headers)
    for call in calls:
        analysis = call.get("analysis", {})
        sheet.append([call.get("call_sid"), call.get("status"), call.get("updated_at"), analysis.get("sentiment"), excel_value(analysis.get("agent_performance_score")), analysis.get("issue_resolved"), analysis.get("revenue_opportunity"), "; ".join(analysis.get("keywords", [])), "; ".join(analysis.get("objections", [])), " ".join(analysis.get("missed_opportunities", [])), " ".join(analysis.get("training_recommendations", []))])
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="0B525B")
    sheet.freeze_panes = "A2"
    for column in sheet.columns:
        sheet.column_dimensions[column[0].column_letter].width = min(max(len(str(cell.value or "")) for cell in column) + 2, 45)
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def run() -> None:
    region = env("AWS_REGION")
    table = boto3.resource("dynamodb", region_name=region).Table(env("CALLS_TABLE"))
    calls = []
    page = table.scan(FilterExpression="#status = :completed", ExpressionAttributeNames={"#status": "status"}, ExpressionAttributeValues={":completed": "completed"})
    calls.extend(page.get("Items", []))
    while page.get("LastEvaluatedKey"):
        page = table.scan(FilterExpression="#status = :completed", ExpressionAttributeNames={"#status": "status"}, ExpressionAttributeValues={":completed": "completed"}, ExclusiveStartKey=page["LastEvaluatedKey"])
        calls.extend(page.get("Items", []))
    key = f"reports/master/call-intelligence-{datetime.now(UTC):%Y-%m-%d}.xlsx"
    bucket = env("CALL_DATA_BUCKET")
    storage = boto3.client("s3", region_name=region)
    storage.put_object(Bucket=bucket, Key=key, Body=master_workbook(calls), ServerSideEncryption="aws:kms", SSEKMSKeyId=env("CALL_DATA_KMS_KEY_ID"), ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    url = storage.generate_presigned_url("get_object", Params={"Bucket": bucket, "Key": key}, ExpiresIn=3600)
    recipients = [address.strip() for address in env("REPORT_RECIPIENTS").split(",") if address.strip()]
    boto3.client("sesv2", region_name=region).send_email(FromEmailAddress=env("REPORT_SENDER"), Destination={"ToAddresses": recipients}, Content={"Simple": {"Subject": {"Data": "Daily Call Intelligence Report"}, "Body": {"Text": {"Data": f"Secure master report link, expiring in 60 minutes: {url}"}}}})
    print(f"Created encrypted master report: {key}")


if __name__ == "__main__":
    run()
